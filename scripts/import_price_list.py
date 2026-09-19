"""
Импорт прайс-листа поставщика (xlsx) в каталог компании (own_products).

Это вход в продуктовый сценарий: компания загружает свою номенклатуру,
и дальше закупщик под каждую позицию тендера подбирает готовый SKU из
этого каталога — с артикулом, ценой и характеристиками, то есть сразу
пригодный для коммерческого предложения.

Колонки распознаются по названиям в строке заголовка, а не по номерам:
у каждого поставщика свой порядок, а жёсткие индексы ломаются на первом
же чужом прайсе.

Оптовые прайсы дают несколько цен по объёму закупки (70к/150к/250к).
В price кладём САМУЮ ВЫСОКУЮ — она соответствует наименьшему объёму,
то есть это консервативная оценка. Занизить цену в расчёте маржи
опаснее, чем завысить: во втором случае просто не пойдём в тендер, в
первом — пойдём в убыток. Полная шкала сохраняется в price_text.

Запуск:
  python -m scripts.import_price_list <файл.xlsx> --company 57 --supplier Matrix --dry
  python -m scripts.import_price_list <файл.xlsx> --company 57 --supplier Matrix
"""
import argparse
import asyncio
import logging
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(os.path.dirname(__file__), "..", ".env"))
except ImportError:
    pass

from sqlalchemy import select

from database import DatabaseSession, OwnProduct

logging.basicConfig(format="%(asctime)s %(levelname)s %(message)s", level=logging.INFO)
log = logging.getLogger("import_price_list")

# Как называются нужные колонки у разных поставщиков. Сопоставление по
# вхождению подстроки в заголовок, в нижнем регистре.
COLUMN_HINTS = {
    'sku': ('артикул', 'код товара', 'sku', 'код'),
    'name': ('описание продукции', 'наименование', 'описание', 'товар'),
    'color': ('цвет',),
    'size': ('размер',),
    'weight': ('вес',),
    'unit': ('ед / измер', 'ед/измер', 'единица', 'ед. изм'),
    'pack': ('упаковка',),
    'box': ('трансп', 'короб'),
    'stock': ('наличие', 'склад'),
}
# Колонка с ценой: заголовок вида «сумма 70 000 руб» или просто «цена».
PRICE_HINTS = ('цена', 'сумма', 'руб', 'стоимость')


def _norm(value) -> str:
    return ' '.join(str(value).split()) if value is not None else ''


def find_header_row(ws, max_scan: int = 15):
    """Строка заголовка — первая, где нашлись и артикул, и наименование."""
    for row_idx in range(1, min(max_scan, ws.max_row) + 1):
        titles = [_norm(c.value).lower() for c in ws[row_idx]]
        joined = ' | '.join(titles)
        if any(h in joined for h in COLUMN_HINTS['sku']) and \
           any(h in joined for h in COLUMN_HINTS['name']):
            return row_idx
    return None


def map_columns(ws, header_row: int):
    """Название колонки -> её номер. Плюс список колонок с ценами."""
    mapping = {}
    price_cols = []
    for idx, cell in enumerate(ws[header_row], start=1):
        title = _norm(cell.value).lower()
        if not title:
            continue
        for key, hints in COLUMN_HINTS.items():
            if key not in mapping and any(h in title for h in hints):
                mapping[key] = idx
                break
        else:
            if any(h in title for h in PRICE_HINTS):
                price_cols.append((idx, _norm(cell.value)))
    return mapping, price_cols


def _to_price(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r'[^\d.,]', '', str(value)).replace(',', '.')
    try:
        return float(text)
    except ValueError:
        return None


def parse_rows(ws, header_row: int, mapping, price_cols):
    """Строки товаров. Разделы и пустые строки пропускаются."""
    items = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=False):
        def cell(key):
            col = mapping.get(key)
            return _norm(row[col - 1].value) if col and col <= len(row) else ''

        sku, name = cell('sku'), cell('name')
        # Заголовок раздела («ПЕРЧАТКИ ЛАТЕКСНЫЕ», «Matrix Premium») —
        # это строка с текстом, но без артикула. Товаром не является.
        if not sku or not name:
            continue

        prices = []
        for col, title in price_cols:
            if col <= len(row):
                p = _to_price(row[col - 1].value)
                if p:
                    prices.append((p, title))
        if not prices:
            continue

        # Консервативно: цена наименьшего объёма, то есть самая высокая.
        base_price = max(p for p, _ in prices)
        price_text = '; '.join(f'{p:g} ({t})' for p, t in prices)[:120]

        params = ', '.join(filter(None, [
            f"цвет {cell('color')}" if cell('color') else '',
            f"вес {cell('weight')}" if cell('weight') else '',
        ]))
        pack = ' / '.join(filter(None, [cell('pack'), cell('box')]))

        items.append({
            'sku': sku[:80],
            'name': name[:300],
            'sizes': cell('size')[:200] or None,
            'params': params or None,
            'pack': pack[:120] or None,
            'price': base_price,
            'price_unit': cell('unit')[:40] or None,
            'price_text': price_text or None,
            'notes': cell('stock')[:200] or None,
        })
    return items


async def save(items, company_id: int, supplier: str, category: str,
               dry: bool) -> None:
    async with DatabaseSession() as session:
        existing = {
            p.sku: p for p in (await session.scalars(
                select(OwnProduct).where(
                    OwnProduct.company_id == company_id,
                    OwnProduct.supplier == supplier,
                )
            )).all() if p.sku
        }
        created = updated = 0
        for item in items:
            row = existing.get(item['sku'])
            if row:
                # Повторный импорт того же прайса — обновление цен, а не
                # дубли: артикул уникален в пределах поставщика.
                for key, value in item.items():
                    setattr(row, key, value)
                updated += 1
            else:
                created += 1
                if not dry:
                    session.add(OwnProduct(
                        company_id=company_id, supplier=supplier,
                        category=category, source=supplier, **item))
        if dry:
            session.expunge_all()
        else:
            await session.commit()
        log.info("%s создано %d, обновлено %d",
                 "DRY, не записано —" if dry else "записано:", created, updated)


async def main(path: str, company_id: int, supplier: str, category: str,
               dry: bool) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    total = 0
    for ws in wb.worksheets:
        header_row = find_header_row(ws)
        if not header_row:
            log.info("лист «%s»: заголовок не найден, пропускаем", ws.title)
            continue
        mapping, price_cols = map_columns(ws, header_row)
        log.info("лист «%s»: заголовок в строке %d, колонки %s, цен %d",
                 ws.title, header_row, sorted(mapping), len(price_cols))
        items = parse_rows(ws, header_row, mapping, price_cols)
        log.info("  распознано позиций: %d", len(items))
        for it in items[:3]:
            log.info("    %s | %s | %s | %s ₽/%s",
                     it['sku'], it['name'][:40], it['sizes'],
                     it['price'], it['price_unit'])
        if items:
            await save(items, company_id, supplier, category, dry)
            total += len(items)
    log.info("итого позиций: %d", total)


def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument('path')
    p.add_argument('--company', type=int, required=True)
    p.add_argument('--supplier', required=True)
    p.add_argument('--category', default='siz')
    p.add_argument('--dry', action='store_true')
    return p.parse_args()


if __name__ == '__main__':
    a = parse_args()
    asyncio.run(main(a.path, a.company, a.supplier, a.category, a.dry))
